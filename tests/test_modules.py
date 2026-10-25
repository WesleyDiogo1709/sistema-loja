import io
import os
import sqlite3
import unittest
from datetime import datetime, timedelta

from tests.helpers import Api, BaseCase, ADMIN_PW

TODAY = datetime.now().strftime("%Y-%m-%d")


class DeviceTests(BaseCase):
    def device(self, **kw):
        d = {"brand": "Apple", "model": "iPhone 13", "storage": "128 GB", "color": "Preto", "imei1": self.valid_imei(),
             "serial": "F2LXK1234", "battery_health": 88, "condition": "Muito bom", "origin": "Compra de cliente",
             "cost_cents": 250000, "price_cents": 320000, **kw}
        return self.admin.ok(self.admin.post("/api/devices", d))

    def test_create_with_individual_code_and_history(self):
        d = self.device()
        self.assertEqual(d["code"], f"APL{d['id']:06d}")
        self.assertEqual(d["status"], "disponivel")
        det = self.admin.ok(self.admin.get(f"/api/devices/{d['id']}"))
        self.assertEqual(det["events"][0]["event"], "entrada")

    def test_imei_validation_and_uniqueness(self):
        self.assertEqual(self.admin.post("/api/devices", {"brand": "Apple", "model": "X", "imei1": "123"})[0], 400)
        self.assertEqual(self.admin.post("/api/devices", {"brand": "Apple", "model": "X", "imei1": "356920112345678"})[0], 400)  # Luhn
        a = self.device()
        s, b = self.admin.post("/api/devices", {"brand": "Apple", "model": "X", "imei1": a["imei1"]})
        self.assertEqual((s, b["code"]), (409, "duplicate_device"))
        s, b = self.admin.post("/api/devices", {"brand": "Apple", "model": "X", "imei2": a["imei1"], "serial": "OUTRA"})
        self.assertEqual(s, 409)
        s, b = self.admin.post("/api/devices", {"brand": "Apple", "model": "X", "serial": "F2LXK1234"})
        self.assertEqual(s, 409)
        self.assertEqual(self.admin.post("/api/devices", {"brand": "Apple", "model": "X", "imei1": self.valid_imei("35692011234568")})[0], 201)

    def test_search_by_imei_serial_code(self):
        d = self.device()
        for q in (d["imei1"], "F2LXK1234", d["code"], d["imei1"][-6:]):
            self.assertEqual([x["id"] for x in self.admin.ok(self.admin.get(f"/api/devices?q={q}"))], [d["id"]], q)

    def test_lookup_by_scanner_returns_device(self):
        d = self.device()
        r = self.admin.ok(self.admin.get(f"/api/lookup?code={d['imei1']}"))
        self.assertEqual((r["found"], r["device"]["id"]), ("device", d["id"]))

    def test_sell_device_blocks_second_sale_and_restores_on_cancel(self):
        d = self.device()
        self.open_cash()
        sale = self.admin.ok(self.sale([{"device_id": d["id"]}], self.cash(320000)))["sale"]
        self.assertEqual(sale["items"][0]["unit_cost_cents"], 250000)
        self.assertEqual(sale["profit_cents"], 70000)
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/devices/{d['id']}"))["status"], "vendido")
        s, b = self.sale([{"device_id": d["id"]}], self.cash(320000))
        self.assertEqual((s, b["code"]), (409, "device_unavailable"))
        self.assertEqual(self.admin.post(f"/api/devices/{d['id']}/status", {"status": "disponivel"})[0], 409)
        self.admin.ok(self.admin.post(f"/api/sales/{sale['id']}/cancel", {"reason": "desistência"}))
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/devices/{d['id']}"))["status"], "disponivel")
        self.admin.ok(self.sale([{"device_id": d["id"]}], self.cash(320000)))  # pode vender de novo

    def test_reserved_or_in_repair_cannot_be_sold(self):
        d = self.device()
        self.open_cash()
        for st in ("reservado", "em_manutencao", "em_teste", "devolvido"):
            self.admin.ok(self.admin.post(f"/api/devices/{d['id']}/status", {"status": st, "reserved_for": "Maria"}))
            s, b = self.sale([{"device_id": d["id"]}], self.cash(320000))
            self.assertEqual((s, b["code"]), (409, "device_unavailable"), st)

    def test_device_warranty_recorded_in_sale(self):
        d = self.device(warranty_days=180)
        self.open_cash()
        sale = self.admin.ok(self.sale([{"device_id": d["id"]}], self.cash(320000)))["sale"]
        self.assertEqual(sale["items"][0]["warranty_days"], 180)
        ends = (datetime.now() + timedelta(days=180)).strftime("%Y-%m-%d")
        self.assertEqual(sale["items"][0]["warranty_ends"], ends)

    def test_cost_hidden_from_seller(self):
        d = self.device()
        v = self.user("vera", "vendedor")
        self.assertEqual(v.post("/api/devices", {"brand": "A", "model": "B"})[0], 403)
        got = v.ok(v.get(f"/api/devices/{d['id']}"))
        self.assertNotIn("cost_cents", got)
        self.assertEqual(got["price_cents"], 320000)


class ServiceOrderTests(BaseCase):
    def order(self, **kw):
        d = {"customer_name": "Carlos", "customer_phone": "(21) 98888-7777", "brand": "Samsung", "model": "A54",
             "imei": self.valid_imei(), "defect": "Tela quebrada", "physical_state": "Riscos na traseira",
             "accessories": "Capa", "labor_cents": 8000, "deadline": (datetime.now() + timedelta(days=3)).strftime("%Y-%m-%d"), **kw}
        return self.admin.ok(self.admin.post("/api/os", d))["id"]

    def test_create_and_validations(self):
        self.assertEqual(self.admin.post("/api/os", {"customer_name": "X", "customer_phone": "1", "brand": "A", "model": "B"})[0], 400)  # sem defeito
        self.assertEqual(self.admin.post("/api/os", {"customer_name": "X", "brand": "A", "model": "B", "defect": "d"})[0], 400)  # sem telefone
        oid = self.order()
        o = self.admin.ok(self.admin.get(f"/api/os/{oid}"))
        self.assertEqual((o["status"], o["total_cents"], o["events"][0]["status"]), ("recebido", 8000, "recebido"))

    def test_full_repair_flow_with_part_and_payments(self):
        tela = self.product(name="Tela A54", cost=15000, price=25000, qty=3)
        tech = self.user("tato", "tecnico")
        oid = self.order()
        # técnico não vê custos
        t = tech.ok(tech.post(f"/api/os/{oid}/parts", {"product_id": tela["id"], "qty": 1}))
        self.assertNotIn("parts_cost_cents", t)
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/products/{tela['id']}"))["qty"], 2)
        o = self.admin.ok(self.admin.get(f"/api/os/{oid}"))
        self.assertEqual((o["parts_cents"], o["parts_cost_cents"], o["total_cents"]), (25000, 15000, 33000))
        # não avança sem aprovação do cliente
        s, b = self.admin.post(f"/api/os/{oid}/status", {"status": "em_reparo"})
        self.assertEqual((s, b["code"]), (409, "approval_required"))
        self.assertEqual(self.admin.post(f"/api/os/{oid}/approval", {"approved": True, "note": ""})[0], 400)
        self.admin.ok(self.admin.post(f"/api/os/{oid}/approval", {"approved": True, "note": "Autorizou por WhatsApp (Carlos)"}))
        for st in ("em_diagnostico", "em_reparo", "em_testes", "pronto_retirada"):
            self.admin.ok(self.admin.post(f"/api/os/{oid}/status", {"status": st}))
        # pagamento parcial em dinheiro exige caixa aberto
        s, b = self.admin.post(f"/api/os/{oid}/payments", {"payments": [{"method": "dinheiro", "amount_cents": 10000}]})
        self.assertEqual((s, b["code"]), (409, "cash_closed"))
        self.open_cash(opening=0)
        self.admin.ok(self.admin.post(f"/api/os/{oid}/payments", {"payments": [{"method": "dinheiro", "amount_cents": 10000, "received_cents": 10000}]}))
        # entrega com saldo pendente exige confirmação
        s, b = self.admin.post(f"/api/os/{oid}/status", {"status": "entregue"})
        self.assertEqual((s, b["code"], b["balance_cents"]), (409, "pending_balance", 23000))
        self.assertEqual(self.admin.post(f"/api/os/{oid}/payments", {"payments": [{"method": "pix", "amount_cents": 99999}]})[0], 400)  # acima do saldo
        self.admin.ok(self.admin.post(f"/api/os/{oid}/payments", {"payments": [{"method": "pix", "amount_cents": 13000}]}))
        o = self.admin.ok(self.admin.post(f"/api/os/{oid}/status", {"status": "entregue", "allow_pending": True}))
        self.assertEqual((o["status"], o["balance_cents"], o["warranty_ends"] is not None), ("entregue", 10000, True))
        # dinheiro da OS entrou no caixa esperado
        self.assertEqual(self.admin.ok(self.admin.get("/api/cash"))["expected_cash_cents"], 10000)
        # saldo pendente aparece em contas a receber
        rec = self.admin.ok(self.admin.get("/api/finance/receivables"))
        self.assertEqual(rec["totals"]["service_pending_cents"], 10000)
        # encerrada: não aceita mais alterações
        self.assertEqual(self.admin.post(f"/api/os/{oid}/parts", {"product_id": tela["id"], "qty": 1})[0], 409)

    def test_remove_part_returns_stock_and_invalidates_approval(self):
        p = self.product(qty=5)
        oid = self.order()
        self.admin.ok(self.admin.post(f"/api/os/{oid}/approval", {"approved": True, "note": "ok verbal"}))
        o = self.admin.ok(self.admin.post(f"/api/os/{oid}/parts", {"product_id": p["id"], "qty": 2}))
        self.assertEqual(o["approved"], 0)  # mudou o valor: precisa autorizar de novo
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/products/{p['id']}"))["qty"], 3)
        o = self.admin.ok(self.admin.delete(f"/api/os/{oid}/parts/{o['parts'][0]['id']}"))
        self.assertEqual(o["parts"], [])
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/products/{p['id']}"))["qty"], 5)

    def test_part_without_stock_blocked(self):
        p = self.product(qty=1)
        oid = self.order()
        self.assertEqual(self.admin.post(f"/api/os/{oid}/parts", {"product_id": p["id"], "qty": 2})[0], 409)

    def test_cancel_returns_parts_and_refunds_cash(self):
        p = self.product(qty=4, price=5000)
        self.open_cash(opening=0)
        oid = self.order(labor_cents=0)
        self.admin.ok(self.admin.post(f"/api/os/{oid}/parts", {"product_id": p["id"], "qty": 2}))
        self.admin.ok(self.admin.post(f"/api/os/{oid}/payments", {"payments": [{"method": "dinheiro", "amount_cents": 3000}]}))
        self.assertEqual(self.admin.ok(self.admin.get("/api/cash"))["expected_cash_cents"], 3000)
        self.assertEqual(self.admin.post(f"/api/os/{oid}/status", {"status": "cancelado"})[0], 400)  # motivo
        self.admin.ok(self.admin.post(f"/api/os/{oid}/status", {"status": "cancelado", "note": "cliente desistiu"}))
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/products/{p['id']}"))["qty"], 4)
        self.assertEqual(self.admin.ok(self.admin.get("/api/cash"))["expected_cash_cents"], 0)

    def test_discount_requires_authorization(self):
        v = self.user("vera", "vendedor")
        oid = self.order()
        s, b = v.put(f"/api/os/{oid}", {"discount_cents": 500})
        self.assertEqual((s, b["code"]), (403, "authorization_required"))
        self.assertEqual(v.put(f"/api/os/{oid}", {"discount_cents": 500, "authorization": {"username": "admin", "password": ADMIN_PW}})[0], 200)

    def test_search_by_number_phone_name_imei_and_history(self):
        imei = self.valid_imei()
        o1 = self.order(imei=imei)
        o2 = self.order(imei=imei, customer_name="Maria Antônia", customer_phone="(31) 3333-4444")
        for q, expect in ((f"OS-{o1}", [o1]), (str(o2), [o2]), ("maria antonia", [o2]), ("3333", [o2]), (imei, [o2, o1])):
            self.assertEqual([x["id"] for x in self.admin.ok(self.admin.get(f"/api/os?q={q}"))], expect, q)
        self.assertEqual([h["id"] for h in self.admin.ok(self.admin.get(f"/api/os/{o2}"))["history"]], [o1])

    def test_unclaimed_report(self):
        oid = self.order(labor_cents=0)
        for st in ("em_reparo", "pronto_retirada"):
            self.admin.ok(self.admin.post(f"/api/os/{oid}/status", {"status": st}))
        self.assertEqual(self.admin.ok(self.admin.get("/api/os?unclaimed=1")), [])
        c = sqlite3.connect(os.path.join(self.tmp, "loja.db"))
        c.execute("UPDATE service_orders SET completed_at=datetime('now','localtime','-45 day') WHERE id=?", (oid,))
        c.commit(); c.close()
        self.assertEqual([x["id"] for x in self.admin.ok(self.admin.get("/api/os?unclaimed=1"))], [oid])

    def test_print_os_page(self):
        oid = self.order(customer_name="Zé <script>")
        html = self.admin.c.get(f"/print/os/{oid}?formato=a4").get_data(as_text=True)   # A4 continua funcionando
        self.assertIn(f"ORDEM DE SERVIÇO Nº {oid}", html)
        self.assertNotIn("<script>alert", html)
        self.assertIn("Zé &lt;script&gt;", html)
        thermal = self.admin.c.get(f"/print/os/{oid}").get_data(as_text=True)         # padrão: térmica 80 mm
        self.assertIn(f"Nº {oid}", thermal)
        self.assertIn("Zé &lt;script&gt;", thermal)

    def test_technician_cannot_see_cash_or_finance(self):
        t = self.user("tato", "tecnico")
        self.assertEqual(t.get("/api/finance/summary")[0], 403)
        self.assertEqual(t.get("/api/os")[0], 200)
        self.assertEqual(t.get("/api/customers")[0], 403)


class ReturnsAndFinanceTests(BaseCase):
    def test_partial_return_restock_and_cash_refund(self):
        p = self.product(qty=10, cost=1000, price=3000)
        self.open_cash(opening=0)
        sale = self.admin.ok(self.sale([{"product_id": p["id"], "qty": 4}], self.cash(12000)))["sale"]
        item = sale["items"][0]
        rid = self.admin.ok(self.admin.post("/api/returns", {"sale_id": sale["id"], "kind": "devolucao", "reason": "não gostou",
                            "items": [{"sale_item_id": item["id"], "qty": 1, "restock": True}], "refund_method": "dinheiro"}))["id"]
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/products/{p['id']}"))["qty"], 7)
        self.assertEqual(self.admin.ok(self.admin.get("/api/cash"))["expected_cash_cents"], 12000 - 3000)
        # não pode devolver além do vendido
        s, _ = self.admin.post("/api/returns", {"sale_id": sale["id"], "kind": "devolucao", "reason": "x1x",
                               "items": [{"sale_item_id": item["id"], "qty": 4}], "refund_method": "dinheiro"})
        self.assertEqual(s, 409)
        # defeituoso: não volta ao estoque, mas fica registrado como avaria
        self.admin.ok(self.admin.post("/api/returns", {"sale_id": sale["id"], "kind": "devolucao", "reason": "defeito de fábrica",
                      "items": [{"sale_item_id": item["id"], "qty": 1, "restock": False, "condition_note": "não liga"}], "refund_method": "pix"}))
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/products/{p['id']}"))["qty"], 7)
        kinds = [m["kind"] for m in self.admin.ok(self.admin.get(f"/api/products/{p['id']}/history"))["movements"]][:3]
        self.assertEqual(kinds, ["avaria", "devolucao", "devolucao"])
        # venda com devolução não pode ser cancelada
        self.assertEqual(self.admin.post(f"/api/sales/{sale['id']}/cancel", {"reason": "tentativa"})[0], 409)
        got = self.admin.ok(self.admin.get(f"/api/sales/{sale['id']}"))
        self.assertEqual((got["refunded_cents"], got["items"][0]["returned_qty"]), (6000, 2))

    def test_return_requires_permission_or_supervisor(self):
        v = self.user("vera", "vendedor")
        p = self.product(qty=5)
        self.open_cash(v)
        sale = self.admin.ok(self.sale([{"product_id": p["id"], "qty": 1}], self.cash(2500), who=v))["sale"]
        payload = {"sale_id": sale["id"], "kind": "devolucao", "reason": "arrependimento",
                   "items": [{"sale_item_id": sale["items"][0]["id"], "qty": 1}], "refund_method": "dinheiro"}
        s, b = v.post("/api/returns", payload)
        self.assertEqual((s, b["code"]), (403, "authorization_required"))
        self.assertEqual(v.post("/api/returns", {**payload, "authorization": {"username": "admin", "password": ADMIN_PW}})[0], 201)

    def test_exchange_without_refund(self):
        p = self.product(qty=5)
        self.open_cash()
        sale = self.admin.ok(self.sale([{"product_id": p["id"], "qty": 1}], self.cash(2500)))["sale"]
        self.admin.ok(self.admin.post("/api/returns", {"sale_id": sale["id"], "kind": "troca", "reason": "tamanho errado", "refund_cents": 0,
                      "items": [{"sale_item_id": sale["items"][0]["id"], "qty": 1}]}))
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/products/{p['id']}"))["qty"], 5)

    def test_return_of_device(self):
        d = self.admin.ok(self.admin.post("/api/devices", {"brand": "Apple", "model": "11", "cost_cents": 100000, "price_cents": 150000}))
        self.open_cash()
        sale = self.admin.ok(self.sale([{"device_id": d["id"]}], self.cash(150000)))["sale"]
        self.admin.ok(self.admin.post("/api/returns", {"sale_id": sale["id"], "kind": "devolucao", "reason": "garantia",
                      "items": [{"sale_item_id": sale["items"][0]["id"], "qty": 1}], "refund_method": "dinheiro"}))
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/devices/{d['id']}"))["status"], "devolvido")

    def test_finance_numbers_end_to_end(self):
        """Cenário com valores conhecidos: faturamento, custo, taxas, despesas e lucro."""
        self.admin.ok(self.admin.put("/api/settings", {"settings": {"default_acquirer": "Stone"}, "card_fees": [
            {"acquirer": "Stone", "method": "credito", "installments": 1, "fee_percent": "4", "days": 30}]}))
        a = self.product(name="A", cost=1000, price=3000, qty=20)   # lucro 2000/un
        b = self.product(name="B", cost=500, price=1500, qty=20)
        self.open_cash(opening=0)
        # Venda 1: 2xA + 1xB = 7500, desconto 500 => 7000; paga no cartão (taxa 4% = 280)
        s1 = self.admin.ok(self.sale([{"product_id": a["id"], "qty": 2, "discount_cents": 500}, {"product_id": b["id"], "qty": 1}],
                                     [{"method": "credito", "amount_cents": 7000, "installments": 1}]))["sale"]
        # Venda 2: 1xA = 3000 em dinheiro
        self.admin.ok(self.sale([{"product_id": a["id"], "qty": 1}], self.cash(3000)))
        # Venda 3 cancelada não conta
        s3 = self.admin.ok(self.sale([{"product_id": b["id"], "qty": 5}], self.cash(7500)))["sale"]
        self.admin.ok(self.admin.post(f"/api/sales/{s3['id']}/cancel", {"reason": "erro de lançamento"}))
        # Devolução de 1xA da venda 2 (3000) com volta ao estoque
        sale2 = self.admin.ok(self.admin.get("/api/sales"))
        sale2_id = [x for x in sale2 if x["total_cents"] == 3000][0]["id"]
        item = self.admin.ok(self.admin.get(f"/api/sales/{sale2_id}"))["items"][0]
        self.admin.ok(self.admin.post("/api/returns", {"sale_id": sale2_id, "kind": "devolucao", "reason": "desistência",
                      "items": [{"sale_item_id": item["id"], "qty": 1, "restock": True}], "refund_method": "dinheiro"}))
        # Despesa fixa 1200 e variável 300 no mês; compra de mercadoria 4000 (não é despesa)
        for kind, desc, val in (("fixa", "Aluguel", 1200), ("variavel", "Embalagens", 300), ("compra", "Reposição", 4000)):
            self.admin.ok(self.admin.post("/api/payables", {"kind": kind, "description": desc, "amount_cents": val, "due_date": TODAY}))
        # Serviço entregue: peça A (preço 3000, custo 1000) + mão de obra 2000 = 5000
        oid = self.admin.ok(self.admin.post("/api/os", {"customer_name": "Ana", "customer_phone": "1", "brand": "X", "model": "Y", "defect": "z", "labor_cents": 2000}))["id"]
        self.admin.ok(self.admin.post(f"/api/os/{oid}/parts", {"product_id": a["id"], "qty": 1}))
        self.admin.ok(self.admin.post(f"/api/os/{oid}/approval", {"approved": True, "note": "balcão"}))
        for st in ("em_reparo", "pronto_retirada"):
            self.admin.ok(self.admin.post(f"/api/os/{oid}/status", {"status": st}))
        self.admin.ok(self.admin.post(f"/api/os/{oid}/payments", {"payments": [{"method": "pix", "amount_cents": 5000}]}))
        self.admin.ok(self.admin.post(f"/api/os/{oid}/status", {"status": "entregue"}))
        # garantia com custo 400
        cid = self.admin.ok(self.admin.post("/api/warranty/claims", {"item_desc": "Peça A", "defect": "falha", "customer_name": "Ana"}))["id"]
        self.admin.ok(self.admin.put(f"/api/warranty/claims/{cid}", {"status": "resolvida", "diagnosis": "defeito", "solution": "troca", "resolution": "troca", "cost_cents": 400}))

        f = self.admin.ok(self.admin.get(f"/api/finance/summary?from={TODAY}&to={TODAY}"))
        self.assertEqual(f["sales_total_cents"], 7000 + 3000)           # vendas concluídas
        self.assertEqual(f["returns_cents"], 3000)
        self.assertEqual(f["services_total_cents"], 5000)
        self.assertEqual(f["revenue_cents"], 10000 - 3000 + 5000)       # faturamento = 12000
        # custo: V1 (2*1000+500=2500) + V2 (1000) - devolvido que voltou (1000) + peça do serviço (1000) = 3500
        self.assertEqual(f["cost_cents"], 3500)
        self.assertEqual(f["gross_profit_cents"], 12000 - 3500)         # 8500
        self.assertEqual(f["card_fees_cents"], 280)
        self.assertEqual((f["expenses_cents"], f["purchases_cents"], f["warranty_cost_cents"]), (1500, 4000, 400))
        self.assertEqual(f["result_cents"], 8500 - 280 - 1500 - 400)    # 6320
        # recebimentos (o que entrou de fato): cartão 7000 + dinheiro 3000 + PIX do serviço 5000
        self.assertEqual(f["received_cents"], 15000)
        self.assertEqual(f["received_by_method"], {"credito": 7000, "dinheiro": 3000, "pix": 5000})
        # o vendedor/gerente sem "ver_lucros" não enxerga custo nem lucro
        roles = self.admin.ok(self.admin.get("/api/permissions"))["roles"]
        roles["gerente"].remove("ver_lucros"); roles["gerente"].remove("ver_custos")
        self.admin.ok(self.admin.put("/api/permissions", {"roles": roles}))
        g = self.user("gabi", "gerente")
        fg = g.ok(g.get(f"/api/finance/summary?from={TODAY}&to={TODAY}"))
        self.assertNotIn("gross_profit_cents", fg); self.assertNotIn("cost_cents", fg)
        self.assertEqual(g.get("/api/reports/lucro_produto")[0], 403)
        # fluxo de caixa: entradas 15000; saídas: devolução em dinheiro 3000
        cf = self.admin.ok(self.admin.get(f"/api/finance/cashflow?from={TODAY}&to={TODAY}"))
        self.assertEqual((cf["in_cents"], cf["out_cents"]), (15000, 3000))

    def test_card_receivables_forecast_net_of_fees(self):
        self.admin.ok(self.admin.put("/api/settings", {"settings": {"default_acquirer": "Stone"}, "card_fees": [
            {"acquirer": "Stone", "method": "credito", "installments": 1, "fee_percent": "3", "days": 30},
            {"acquirer": "Stone", "method": "credito", "installments": 3, "fee_percent": "6", "days": 30}]}))
        p = self.product(price=30000)
        self.open_cash()
        self.admin.ok(self.sale([{"product_id": p["id"], "qty": 1}], [{"method": "credito", "amount_cents": 30000, "installments": 3}]))
        cards = self.admin.ok(self.admin.get("/api/finance/receivables"))["cards"]
        self.assertEqual([c["amount_cents"] for c in cards], [9400, 9400, 9400])   # 30000 - 6% = 28200 / 3
        self.assertEqual([c["due_date"] for c in cards], [(datetime.now() + timedelta(days=d)).strftime("%Y-%m-%d") for d in (30, 60, 90)])

    def test_payables_flow(self):
        ids = self.admin.ok(self.admin.post("/api/payables", {"kind": "fixa", "description": "Aluguel", "amount_cents": 150000,
                            "due_date": "2026-01-31", "repeat_months": 3}))["ids"]
        due = [p["due_date"] for p in self.admin.ok(self.admin.get("/api/payables?status=abertas"))]
        self.assertEqual(due, ["2026-01-31", "2026-02-28", "2026-03-31"])
        self.assertTrue(all(p["overdue"] for p in self.admin.ok(self.admin.get("/api/payables"))))
        self.admin.ok(self.admin.post(f"/api/payables/{ids[0]}/pay", {"method": "pix"}))
        self.assertEqual(self.admin.post(f"/api/payables/{ids[0]}/pay", {})[0], 409)
        self.assertEqual(self.admin.delete(f"/api/payables/{ids[0]}")[0], 409)   # paga: histórico
        self.assertEqual(self.admin.delete(f"/api/payables/{ids[1]}")[0], 200)

    def test_warranty_lookup_and_claim(self):
        d = self.admin.ok(self.admin.post("/api/devices", {"brand": "Apple", "model": "12", "imei1": self.valid_imei(), "price_cents": 200000, "cost_cents": 100000}))
        c = self.admin.ok(self.admin.post("/api/customers", {"name": "Paula", "phone": "11988887777"}))
        self.open_cash()
        sale = self.admin.ok(self.sale([{"device_id": d["id"]}], self.cash(200000), customer_id=c["id"]))["sale"]
        res = self.admin.ok(self.admin.get(f"/api/warranty/lookup?q={d['imei1']}"))
        self.assertEqual((res[0]["sale_id"], res[0]["in_warranty"]), (sale["id"], True))
        self.assertEqual(self.admin.ok(self.admin.get("/api/warranty/lookup?q=11988887777"))[0]["customer_name"], "Paula")
        cid = self.admin.ok(self.admin.post("/api/warranty/claims", {"sale_item_id": res[0]["sale_item_id"], "defect": "bateria viciada"}))["id"]
        # não encerra sem diagnóstico e solução
        self.assertEqual(self.admin.put(f"/api/warranty/claims/{cid}", {"status": "resolvida"})[0], 400)
        self.admin.ok(self.admin.put(f"/api/warranty/claims/{cid}", {"status": "em_analise", "diagnosis": "bateria 70%"}))
        self.admin.ok(self.admin.put(f"/api/warranty/claims/{cid}", {"status": "resolvida", "solution": "troca de bateria", "resolution": "reparo", "cost_cents": 12000}))
        cl = self.admin.ok(self.admin.get("/api/warranty/claims"))[0]
        self.assertEqual((cl["status"], cl["customer_name"], cl["cost_cents"]), ("resolvida", "Paula", 12000))
        self.assertEqual(self.admin.put(f"/api/warranty/claims/{cid}", {"status": "aberta"})[0], 409)


class ReportsBackupTests(BaseCase):
    def seed(self):
        p = self.product(name="Fone Bluetooth", barcode="7891000100103", cost=2000, price=5000, qty=10, category="Áudio")
        self.open_cash()
        for _ in range(2):
            self.admin.ok(self.sale([{"product_id": p["id"], "qty": 1}], self.cash(5000)))
        return p

    def test_reports_json_csv_xlsx(self):
        self.seed()
        rep = self.admin.ok(self.admin.get(f"/api/reports/lucro_produto?from={TODAY}&to={TODAY}"))
        row = rep["rows"][0]
        self.assertEqual((row["qty"], row["revenue_cents"], row["cost_cents"], row["profit_cents"], row["margin"]), (2, 10000, 4000, 6000, 60.0))
        status, csv = self.admin.get("/api/reports/vendas?format=csv")
        text = csv.decode("utf-8")
        self.assertTrue(text.startswith("﻿"))
        self.assertIn("50,00", text)
        status, xlsx = self.admin.get("/api/reports/mais_vendidos?format=xlsx")
        self.assertEqual(status, 200)
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(xlsx)).active
        self.assertEqual(ws["A4"].value, "Produto")
        self.assertEqual(ws["C5"].value, 2)
        self.assertEqual(ws["D5"].value, 100.0)

    def test_report_filters(self):
        self.seed()
        self.assertEqual(len(self.admin.ok(self.admin.get("/api/reports/vendas?category=Áudio"))["rows"]), 2)
        self.assertEqual(len(self.admin.ok(self.admin.get("/api/reports/vendas?category=Outra"))["rows"]), 0)
        self.assertEqual(len(self.admin.ok(self.admin.get("/api/reports/vendas?method=pix"))["rows"]), 0)
        self.assertEqual(len(self.admin.ok(self.admin.get("/api/reports/vendas?method=dinheiro"))["rows"]), 2)
        self.assertEqual(len(self.admin.ok(self.admin.get("/api/reports/vendas?from=2000-01-01&to=2000-01-02"))["rows"]), 0)

    def test_csv_formula_injection_neutralized(self):
        self.product(name='=HYPERLINK("http://x")', qty=1)
        text = self.admin.get("/api/reports/estoque?format=csv")[1].decode("utf-8")
        self.assertNotIn(';=HYPERLINK', text)
        self.assertIn("'=HYPERLINK", text)

    def test_stalled_report(self):
        p = self.product(name="Encalhado", qty=3)
        c = sqlite3.connect(os.path.join(self.tmp, "loja.db"))
        c.execute("UPDATE products SET created_at=datetime('now','localtime','-90 day') WHERE id=?", (p["id"],))
        c.commit(); c.close()
        rep = self.admin.ok(self.admin.get("/api/reports/parados"))
        self.assertEqual([r["name"] for r in rep["rows"]], ["Encalhado"])

    def test_backup_and_restore_roundtrip(self):
        p = self.seed()
        name = self.admin.ok(self.admin.post("/api/backups"))["name"]
        # altera dados depois do backup
        self.admin.ok(self.sale([{"product_id": p["id"], "qty": 3}], self.cash(15000)))
        self.admin.ok(self.admin.post("/api/products", {"name": "Novo depois", "price_cents": 1}))
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/products/{p['id']}"))["qty"], 5)
        # confirmação e senha são exigidos
        self.assertEqual(self.admin.req("POST", "/api/backups/restore", files={"name": name, "confirm": "sim", "password": ADMIN_PW})[0], 400)
        self.assertEqual(self.admin.req("POST", "/api/backups/restore", files={"name": name, "confirm": "RESTAURAR", "password": "errada"})[0], 403)
        res = self.admin.ok(self.admin.req("POST", "/api/backups/restore", files={"name": name, "confirm": "RESTAURAR", "password": ADMIN_PW}))
        self.assertTrue(res["safety_backup"].startswith("loja-antes-restauracao"))
        # sessões são derrubadas; novo login e dados voltam ao estado do backup
        self.assertEqual(self.admin.get("/api/products")[0], 401)
        self.admin.ok(self.admin.post("/api/login", {"username": "admin", "password": ADMIN_PW}))
        self.assertEqual(self.admin.ok(self.admin.get(f"/api/products/{p['id']}"))["qty"], 8)
        self.assertEqual([x["name"] for x in self.admin.ok(self.admin.get("/api/products?q=Novo"))], [])
        self.assertEqual(len(self.admin.ok(self.admin.get("/api/sales"))), 2)

    def test_restore_from_uploaded_file_and_rejects_invalid(self):
        self.seed()
        name = self.admin.ok(self.admin.post("/api/backups"))["name"]
        status, data = self.admin.get(f"/api/backups/{name}/download")
        self.assertEqual(status, 200)
        self.assertTrue(data.startswith(b"SQLite format 3"))
        bad = self.admin.req("POST", "/api/backups/restore", files={"file": (io.BytesIO(b"isto nao e um banco"), "x.db"), "confirm": "RESTAURAR", "password": ADMIN_PW})
        self.assertEqual(bad[0], 400)
        other = sqlite3.connect(os.path.join(self.tmp, "outro.db")); other.execute("create table t(x)"); other.commit(); other.close()
        with open(os.path.join(self.tmp, "outro.db"), "rb") as fh:
            bad = self.admin.req("POST", "/api/backups/restore", files={"file": (io.BytesIO(fh.read()), "outro.db"), "confirm": "RESTAURAR", "password": ADMIN_PW})
        self.assertEqual(bad[0], 400)
        self.assertIn("não é um backup", bad[1]["error"])
        ok = self.admin.req("POST", "/api/backups/restore", files={"file": (io.BytesIO(data), "b.db"), "confirm": "RESTAURAR", "password": ADMIN_PW})
        self.assertEqual(ok[0], 200)

    def test_backup_path_traversal_blocked(self):
        self.assertEqual(self.admin.get("/api/backups/..%2F..%2Floja.db/download")[0], 404)
        self.assertEqual(self.admin.get("/api/backups/loja-x-1.db/download")[0], 404)

    def test_data_persists_across_restart(self):
        self.seed()
        from loja.app import create_app
        app2 = create_app(testing=True)
        c = Api(app2)
        c.ok(c.post("/api/login", {"username": "admin", "password": ADMIN_PW}))
        self.assertEqual(len(c.ok(c.get("/api/sales"))), 2)
        self.assertEqual(c.ok(c.get("/api/products"))[0]["qty"], 8)

    def test_audit_log_records_sensitive_actions(self):
        p = self.seed()
        self.admin.ok(self.admin.post(f"/api/products/{p['id']}/adjust", {"new_qty": 1, "reason": "teste de auditoria"}))
        log = self.admin.ok(self.admin.get("/api/audit"))
        actions = {r["action"] for r in log}
        self.assertTrue({"login", "venda_criada", "ajuste_estoque", "caixa_aberto", "produto_criado"} <= actions)
        adj = [r for r in log if r["action"] == "ajuste_estoque"][0]
        self.assertIn('"qty": 8', adj["before_json"]); self.assertIn('"qty": 1', adj["after_json"])
        self.assertEqual(adj["user_name"], "Administrador")

    def test_settings_payment_methods_protected(self):
        self.seed()
        s = self.admin.ok(self.admin.get("/api/settings"))
        methods = [m for m in self.admin.ok(self.admin.get("/api/me"))["payment_methods"] if m["key"] != "dinheiro"]
        self.assertEqual(self.admin.put("/api/settings", {"payment_methods": methods})[0], 400)   # sem dinheiro
        methods = [m for m in self.admin.ok(self.admin.get("/api/me"))["payment_methods"]] + [{"label": "Vale-troca", "type": "other"}]
        self.admin.ok(self.admin.put("/api/settings", {"payment_methods": methods}))
        self.assertIn("valetroca", [m["key"] for m in self.admin.ok(self.admin.get("/api/me"))["payment_methods"]])
        used_removed = [m for m in self.admin.ok(self.admin.get("/api/me"))["payment_methods"] if m["key"] != "dinheiro"] + [{"label": "Dinheiro2", "type": "cash"}]
        self.assertEqual(self.admin.put("/api/settings", {"payment_methods": used_removed})[0], 400)  # "dinheiro" já foi usado


if __name__ == "__main__":
    unittest.main()
