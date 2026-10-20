"""Relatórios com filtros e exportação (CSV, Excel). PDF: tela de impressão do navegador ("Salvar como PDF")."""
import csv
import io
from datetime import datetime

from .services import ApiError, now, permissions_of

COL = lambda key, label, type="text": {"key": key, "label": label, "type": type}


def _range(f):
    a = f.get("from") or "1970-01-01"
    b = f.get("to") or "2999-12-31"
    return a, b


def r_vendas(conn, f, perms):
    a, b = _range(f)
    sql = ("SELECT s.id, s.created_at, u.name AS seller, COALESCE(c.name,'') AS customer, s.status, s.subtotal_cents, s.discount_cents, s.total_cents,"
           " (SELECT GROUP_CONCAT(method || CASE WHEN installments>1 THEN ' ' || installments || 'x' ELSE '' END, ', ') FROM sale_payments WHERE sale_id=s.id) AS methods,"
           " (SELECT COALESCE(SUM(unit_cost_cents*qty),0) FROM sale_items WHERE sale_id=s.id) AS cost_cents,"
           " (SELECT COALESCE(SUM(fee_cents),0) FROM sale_payments WHERE sale_id=s.id) AS fee_cents"
           " FROM sales s JOIN users u ON u.id=s.user_id LEFT JOIN customers c ON c.id=s.customer_id WHERE date(s.created_at) BETWEEN ? AND ?")
    args = [a, b]
    if f.get("user_id"):
        sql += " AND s.user_id=?"; args.append(f["user_id"])
    if f.get("method"):
        sql += " AND EXISTS (SELECT 1 FROM sale_payments p WHERE p.sale_id=s.id AND p.method=?)"; args.append(f["method"])
    if f.get("product_id"):
        sql += " AND EXISTS (SELECT 1 FROM sale_items i WHERE i.sale_id=s.id AND i.product_id=?)"; args.append(f["product_id"])
    if f.get("category"):
        sql += " AND EXISTS (SELECT 1 FROM sale_items i JOIN products p ON p.id=i.product_id WHERE i.sale_id=s.id AND p.category=?)"; args.append(f["category"])
    if f.get("supplier"):
        sql += " AND EXISTS (SELECT 1 FROM sale_items i JOIN products p ON p.id=i.product_id WHERE i.sale_id=s.id AND p.supplier=?)"; args.append(f["supplier"])
    rows = [dict(r) for r in conn.execute(sql + " ORDER BY s.id DESC", args)]
    cols = [COL("id", "Venda", "int"), COL("created_at", "Data", "datetime"), COL("seller", "Vendedor"), COL("customer", "Cliente"),
            COL("status", "Situação"), COL("methods", "Pagamento"), COL("subtotal_cents", "Subtotal", "money"),
            COL("discount_cents", "Desconto", "money"), COL("total_cents", "Total", "money")]
    if "ver_lucros" in perms:
        for r in rows:
            r["profit_cents"] = r["total_cents"] - r["cost_cents"] - r["fee_cents"] if r["status"] == "concluida" else 0
        cols += [COL("cost_cents", "Custo", "money"), COL("fee_cents", "Taxas", "money"), COL("profit_cents", "Lucro estimado", "money")]
    return "Vendas", cols, rows


def r_lucro_produto(conn, f, perms):
    a, b = _range(f)
    sql = ("SELECT COALESCE(p.name, i.product_name) AS name, COALESCE(p.category,'') AS category, SUM(i.qty) AS qty, SUM(i.total_cents) AS revenue_cents,"
           " SUM(i.unit_cost_cents*i.qty) AS cost_cents, SUM(i.total_cents - i.unit_cost_cents*i.qty) AS profit_cents"
           " FROM sale_items i JOIN sales s ON s.id=i.sale_id LEFT JOIN products p ON p.id=i.product_id"
           " WHERE s.status='concluida' AND date(s.created_at) BETWEEN ? AND ?")
    args = [a, b]
    if f.get("category"):
        sql += " AND p.category=?"; args.append(f["category"])
    if f.get("supplier"):
        sql += " AND p.supplier=?"; args.append(f["supplier"])
    if f.get("product_id"):
        sql += " AND i.product_id=?"; args.append(f["product_id"])
    if f.get("user_id"):
        sql += " AND s.user_id=?"; args.append(f["user_id"])
    rows = [dict(r) for r in conn.execute(sql + " GROUP BY COALESCE(i.product_id, -i.device_id) ORDER BY profit_cents DESC", args)]
    for r in rows:
        r["margin"] = round(r["profit_cents"] * 100 / r["revenue_cents"], 1) if r["revenue_cents"] else 0
    cols = [COL("name", "Produto"), COL("category", "Categoria"), COL("qty", "Qtd.", "int"), COL("revenue_cents", "Receita", "money"),
            COL("cost_cents", "Custo", "money"), COL("profit_cents", "Lucro bruto", "money"), COL("margin", "Margem %", "percent")]
    return "Lucro bruto por produto (já considera descontos; não considera taxas de cartão)", cols, rows


def r_lucro_venda(conn, f, perms):
    title, cols, rows = r_vendas(conn, f, perms)
    rows = [r for r in rows if r["status"] == "concluida"]
    keep = ["id", "created_at", "seller", "customer", "total_cents", "cost_cents", "fee_cents", "profit_cents"]
    cols = [c for c in cols if c["key"] in keep]
    for r in rows:
        r["margin"] = round(r["profit_cents"] * 100 / r["total_cents"], 1) if r["total_cents"] else 0
    cols.append(COL("margin", "Margem %", "percent"))
    return "Lucro por venda (total - custo real - taxas de cartão estimadas)", cols, rows


def r_mais_vendidos(conn, f, perms):
    a, b = _range(f)
    sql = ("SELECT COALESCE(p.name, i.product_name) AS name, COALESCE(p.category,'') AS category, SUM(i.qty) AS qty, SUM(i.total_cents) AS revenue_cents"
           " FROM sale_items i JOIN sales s ON s.id=i.sale_id LEFT JOIN products p ON p.id=i.product_id"
           " WHERE s.status='concluida' AND date(s.created_at) BETWEEN ? AND ?")
    args = [a, b]
    if f.get("category"):
        sql += " AND p.category=?"; args.append(f["category"])
    if f.get("supplier"):
        sql += " AND p.supplier=?"; args.append(f["supplier"])
    rows = [dict(r) for r in conn.execute(sql + " GROUP BY COALESCE(i.product_id, -i.device_id) ORDER BY qty DESC, revenue_cents DESC LIMIT 100", args)]
    return "Produtos mais vendidos", [COL("name", "Produto"), COL("category", "Categoria"), COL("qty", "Qtd. vendida", "int"),
                                      COL("revenue_cents", "Receita", "money")], rows


def r_parados(conn, f, perms):
    days = int(f.get("days") or conn.execute("SELECT value FROM settings WHERE key='inactive_days'").fetchone()[0] or 60)
    sql = ("SELECT p.name, COALESCE(p.category,'') AS category, p.qty, p.cost_cents, p.qty*p.cost_cents AS stock_value_cents,"
           " (SELECT MAX(s.created_at) FROM sale_items i JOIN sales s ON s.id=i.sale_id WHERE i.product_id=p.id AND s.status='concluida') AS last_sale"
           " FROM products p WHERE p.active=1 AND p.qty>0 AND p.created_at < datetime('now','localtime',?)"
           " AND NOT EXISTS (SELECT 1 FROM sale_items i JOIN sales s ON s.id=i.sale_id WHERE i.product_id=p.id AND s.status='concluida'"
           " AND s.created_at >= datetime('now','localtime',?))")
    args = [f"-{days} day", f"-{days} day"]
    if f.get("category"):
        sql += " AND p.category=?"; args.append(f["category"])
    if f.get("supplier"):
        sql += " AND p.supplier=?"; args.append(f["supplier"])
    rows = [dict(r) for r in conn.execute(sql + " ORDER BY stock_value_cents DESC", args)]
    cols = [COL("name", "Produto"), COL("category", "Categoria"), COL("qty", "Estoque", "int"), COL("last_sale", "Última venda", "datetime")]
    if "ver_custos" in perms:
        cols += [COL("cost_cents", "Custo unit.", "money"), COL("stock_value_cents", "Valor parado", "money")]
    return f"Produtos parados (sem venda há {days} dias)", cols, rows


def r_estoque(conn, f, perms):
    sql = "SELECT name, COALESCE(category,'') category, COALESCE(brand,'') brand, sku, barcode, qty, min_qty, COALESCE(location,'') location, cost_cents, price_cents, qty*cost_cents AS stock_value_cents FROM products WHERE active=1"
    args = []
    if f.get("category"):
        sql += " AND category=?"; args.append(f["category"])
    if f.get("supplier"):
        sql += " AND supplier=?"; args.append(f["supplier"])
    if f.get("low"):
        sql += " AND qty<=min_qty"
    rows = [dict(r) for r in conn.execute(sql + " ORDER BY name", args)]
    cols = [COL("name", "Produto"), COL("category", "Categoria"), COL("sku", "SKU"), COL("barcode", "Código de barras"),
            COL("qty", "Qtd.", "int"), COL("min_qty", "Mínimo", "int"), COL("location", "Local"), COL("price_cents", "Preço", "money")]
    if "ver_custos" in perms:
        cols += [COL("cost_cents", "Custo", "money"), COL("stock_value_cents", "Valor em estoque", "money")]
    return ("Posição de estoque" + (" (abaixo do mínimo)" if f.get("low") else "")), cols, rows


def r_forma_pagamento(conn, f, perms):
    a, b = _range(f)
    rows = [dict(r) for r in conn.execute(
        "SELECT p.method, COUNT(*) AS n, SUM(p.amount_cents) AS total_cents, SUM(p.fee_cents) AS fee_cents, SUM(p.net_cents) AS net_cents"
        " FROM sale_payments p JOIN sales s ON s.id=p.sale_id WHERE s.status='concluida' AND date(s.created_at) BETWEEN ? AND ?"
        " GROUP BY p.method ORDER BY total_cents DESC", (a, b))]
    return "Vendas por forma de pagamento", [COL("method", "Forma"), COL("n", "Lançamentos", "int"), COL("total_cents", "Total", "money"),
                                             COL("fee_cents", "Taxas estimadas", "money"), COL("net_cents", "Líquido estimado", "money")], rows


def r_vendedor(conn, f, perms):
    a, b = _range(f)
    rows = [dict(r) for r in conn.execute(
        "SELECT u.name AS seller, COUNT(*) n, SUM(s.total_cents) total_cents, SUM(s.discount_cents) discount_cents,"
        " SUM(s.total_cents)/COUNT(*) AS ticket_cents FROM sales s JOIN users u ON u.id=s.user_id"
        " WHERE s.status='concluida' AND date(s.created_at) BETWEEN ? AND ? GROUP BY s.user_id ORDER BY total_cents DESC", (a, b))]
    return "Vendas por vendedor", [COL("seller", "Vendedor"), COL("n", "Vendas", "int"), COL("total_cents", "Total", "money"),
                                   COL("discount_cents", "Descontos", "money"), COL("ticket_cents", "Ticket médio", "money")], rows


def r_perdas(conn, f, perms):
    a, b = _range(f)
    rows = [dict(r) for r in conn.execute(
        "SELECT m.created_at, p.name, m.kind, -m.qty_delta AS qty, COALESCE(m.unit_cost_cents,0)*-m.qty_delta AS loss_cents, m.reason, u.name AS user_name"
        " FROM stock_movements m JOIN products p ON p.id=m.product_id LEFT JOIN users u ON u.id=m.user_id"
        " WHERE m.kind IN ('perda','avaria','troca') AND date(m.created_at) BETWEEN ? AND ? ORDER BY m.id DESC", (a, b))]
    cols = [COL("created_at", "Data", "datetime"), COL("name", "Produto"), COL("kind", "Tipo"), COL("qty", "Qtd.", "int")]
    if "ver_custos" in perms:
        cols.append(COL("loss_cents", "Custo", "money"))
    cols += [COL("reason", "Justificativa"), COL("user_name", "Responsável")]
    return "Perdas, avarias e produtos para troca", cols, rows


def r_movimentacoes(conn, f, perms):
    a, b = _range(f)
    sql = ("SELECT m.created_at, p.name, m.kind, m.qty_delta, m.qty_before, m.qty_after, m.reason, u.name AS user_name FROM stock_movements m"
           " JOIN products p ON p.id=m.product_id LEFT JOIN users u ON u.id=m.user_id WHERE date(m.created_at) BETWEEN ? AND ?")
    args = [a, b]
    if f.get("product_id"):
        sql += " AND m.product_id=?"; args.append(f["product_id"])
    rows = [dict(r) for r in conn.execute(sql + " ORDER BY m.id DESC LIMIT 5000", args)]
    return "Movimentações de estoque", [COL("created_at", "Data", "datetime"), COL("name", "Produto"), COL("kind", "Tipo"), COL("qty_delta", "Qtd.", "int"),
                                        COL("qty_before", "Antes", "int"), COL("qty_after", "Depois", "int"), COL("reason", "Motivo"), COL("user_name", "Usuário")], rows


def r_contas_pagar(conn, f, perms):
    a, b = _range(f)
    rows = [dict(r) for r in conn.execute(
        "SELECT due_date, kind, description, COALESCE(supplier,'') supplier, amount_cents, paid_at, paid_cents FROM payables"
        " WHERE due_date BETWEEN ? AND ? ORDER BY due_date", (a, b))]
    return "Contas a pagar", [COL("due_date", "Vencimento", "date"), COL("kind", "Tipo"), COL("description", "Descrição"), COL("supplier", "Fornecedor"),
                              COL("amount_cents", "Valor", "money"), COL("paid_at", "Pago em", "date"), COL("paid_cents", "Valor pago", "money")], rows


def r_os(conn, f, perms):
    from . import service_orders as so
    a, b = _range(f)
    rows = []
    for o in so.search(conn, limit=5000):
        if a <= o["entry_at"][:10] <= b and (not f.get("user_id") or o["technician_id"] == int(f["user_id"])):
            rows.append(o)
    return "Ordens de serviço", [COL("id", "OS", "int"), COL("entry_at", "Entrada", "datetime"), COL("customer_name", "Cliente"), COL("brand", "Marca"),
                                 COL("model", "Modelo"), COL("status_label", "Situação"), COL("technician_name", "Técnico"),
                                 COL("total_cents", "Total", "money"), COL("paid_cents", "Pago", "money"), COL("balance_cents", "Saldo", "money")], rows


def r_garantias(conn, f, perms):
    a, b = _range(f)
    rows = [dict(r) for r in conn.execute(
        "SELECT opened_at, item_desc, COALESCE(customer_name,'') customer_name, defect, status, resolution, cost_cents FROM warranty_claims"
        " WHERE date(opened_at) BETWEEN ? AND ? ORDER BY id DESC", (a, b))]
    return "Reclamações de garantia", [COL("opened_at", "Abertura", "datetime"), COL("item_desc", "Item"), COL("customer_name", "Cliente"),
                                       COL("defect", "Defeito"), COL("status", "Situação"), COL("resolution", "Resolução"), COL("cost_cents", "Custo", "money")], rows


def r_caixas(conn, f, perms):
    a, b = _range(f)
    rows = [dict(r) for r in conn.execute(
        "SELECT c.id, c.opened_at, c.closed_at, u.name AS opener, c.opening_cents, c.expected_cents, c.counted_cents,"
        " c.counted_cents - c.expected_cents AS diff_cents, COALESCE(c.notes,'') notes FROM cash_sessions c JOIN users u ON u.id=c.opened_by"
        " WHERE date(c.opened_at) BETWEEN ? AND ? ORDER BY c.id DESC", (a, b))]
    return "Fechamentos de caixa", [COL("id", "Caixa", "int"), COL("opened_at", "Abertura", "datetime"), COL("closed_at", "Fechamento", "datetime"),
                                    COL("opener", "Operador"), COL("opening_cents", "Troco inicial", "money"), COL("expected_cents", "Esperado", "money"),
                                    COL("counted_cents", "Contado", "money"), COL("diff_cents", "Diferença", "money"), COL("notes", "Observação")], rows


def r_devolucoes(conn, f, perms):
    a, b = _range(f)
    rows = [dict(r) for r in conn.execute(
        "SELECT r.created_at, r.sale_id, r.kind, r.reason, r.refund_cents, u.name AS user_name FROM returns r JOIN users u ON u.id=r.user_id"
        " WHERE date(r.created_at) BETWEEN ? AND ? ORDER BY r.id DESC", (a, b))]
    return "Devoluções e trocas", [COL("created_at", "Data", "datetime"), COL("sale_id", "Venda", "int"), COL("kind", "Tipo"), COL("reason", "Motivo"),
                                   COL("refund_cents", "Valor devolvido", "money"), COL("user_name", "Responsável")], rows


def r_aparelhos(conn, f, perms):
    rows = [dict(r) for r in conn.execute(
        "SELECT code, brand, model, COALESCE(storage,'') storage, COALESCE(color,'') color, imei1, battery_health, status, cost_cents, price_cents, entry_date FROM devices ORDER BY id DESC")]
    cols = [COL("code", "Código"), COL("brand", "Marca"), COL("model", "Modelo"), COL("storage", "Capacidade"), COL("color", "Cor"), COL("imei1", "IMEI 1"),
            COL("battery_health", "Bateria %", "int"), COL("status", "Situação"), COL("price_cents", "Preço", "money"), COL("entry_date", "Entrada", "date")]
    if "ver_custos" in perms:
        cols.insert(8, COL("cost_cents", "Custo", "money"))
    return "Aparelhos", cols, rows


# nome: (função, permissão exigida além de ver_relatorios)
REPORTS = {
    "vendas": (r_vendas, None), "lucro_produto": (r_lucro_produto, "ver_lucros"), "lucro_venda": (r_lucro_venda, "ver_lucros"),
    "mais_vendidos": (r_mais_vendidos, None), "parados": (r_parados, None), "estoque": (r_estoque, None),
    "forma_pagamento": (r_forma_pagamento, "ver_financeiro"), "vendedor": (r_vendedor, "ver_financeiro"),
    "perdas": (r_perdas, None), "movimentacoes": (r_movimentacoes, None), "contas_pagar": (r_contas_pagar, "ver_financeiro"),
    "os": (r_os, None), "garantias": (r_garantias, None), "caixas": (r_caixas, "ver_financeiro"),
    "devolucoes": (r_devolucoes, None), "aparelhos": (r_aparelhos, None),
}
LABELS = {
    "vendas": "Vendas", "lucro_produto": "Lucro bruto por produto", "lucro_venda": "Lucro por venda", "mais_vendidos": "Mais vendidos",
    "parados": "Produtos parados", "estoque": "Posição de estoque", "forma_pagamento": "Por forma de pagamento", "vendedor": "Por vendedor",
    "perdas": "Perdas e avarias", "movimentacoes": "Movimentações de estoque", "contas_pagar": "Contas a pagar", "os": "Ordens de serviço",
    "garantias": "Garantias", "caixas": "Fechamentos de caixa", "devolucoes": "Devoluções e trocas", "aparelhos": "Aparelhos",
}


def available(perms):
    if "ver_relatorios" not in perms:
        return []
    return [{"key": k, "label": LABELS[k]} for k, (_, need) in REPORTS.items() if not need or need in perms]


def run(conn, name, filters, perms):
    if name not in REPORTS:
        raise ApiError("Relatório não encontrado.", 404)
    fn, need = REPORTS[name]
    if "ver_relatorios" not in perms or (need and need not in perms):
        raise ApiError("Sem permissão para este relatório.", 403)
    title, cols, rows = fn(conn, filters, perms)
    keys = {c["key"] for c in cols}
    rows = [{k: v for k, v in r.items() if k in keys} for r in rows]
    totals = {c["key"]: sum((r.get(c["key"]) or 0) for r in rows) for c in cols if c["type"] in ("money", "int") and c["key"] not in ("id", "sale_id", "battery_health") and rows}
    return {"title": title, "columns": cols, "rows": rows, "totals": totals, "filters": filters, "generated_at": now()}


# ------------------------------------------------------------ exportação

def _fmt_cell(v, t):
    if v is None:
        return ""
    if t == "money":
        return f"{v / 100:.2f}".replace(".", ",")
    if t == "date":
        return datetime.strptime(v[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
    if t == "datetime":
        return datetime.strptime(v[:19], "%Y-%m-%d %H:%M:%S").strftime("%d/%m/%Y %H:%M")
    if t == "percent":
        return f"{v}".replace(".", ",")
    return str(v)


def _safe(text):
    """Evita injeção de fórmula ao abrir CSV/Excel (células iniciadas por = + - @)."""
    if isinstance(text, str) and text[:1] in ("=", "+", "-", "@") and not text.replace(",", "").replace(".", "").replace("-", "").isdigit():
        return "'" + text
    return text


def to_csv(rep):
    out = io.StringIO()
    w = csv.writer(out, delimiter=";", lineterminator="\r\n")
    w.writerow([rep["title"]])
    w.writerow([c["label"] for c in rep["columns"]])
    for r in rep["rows"]:
        w.writerow([_safe(_fmt_cell(r.get(c["key"]), c["type"])) for c in rep["columns"]])
    if rep["totals"]:
        w.writerow(["TOTAL" if i == 0 else _fmt_cell(rep["totals"].get(c["key"]), c["type"]) if c["key"] in rep["totals"] else ""
                    for i, c in enumerate(rep["columns"])])
    return ("﻿" + out.getvalue()).encode("utf-8")  # BOM: o Excel abre com acentos corretos


def to_xlsx(rep):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = rep["title"][:31].replace("/", "-")
    ws.append([rep["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    ws.append([])
    ws.append([c["label"] for c in rep["columns"]])
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(horizontal="center")
    for r in rep["rows"]:
        row = []
        for c in rep["columns"]:
            v = r.get(c["key"])
            if v is None:
                row.append(None)
            elif c["type"] == "money":
                row.append(v / 100)
            elif c["type"] in ("date", "datetime"):
                row.append(datetime.strptime(v[:19] if c["type"] == "datetime" else v[:10], "%Y-%m-%d %H:%M:%S" if c["type"] == "datetime" else "%Y-%m-%d"))
            else:
                row.append(_safe(v))
        ws.append(row)
    for ci, c in enumerate(rep["columns"], 1):
        fmt = {"money": '"R$" #,##0.00', "date": "DD/MM/YYYY", "datetime": "DD/MM/YYYY HH:MM", "percent": '0.0"%"'}.get(c["type"])
        longest = len(c["label"])
        for row in ws.iter_rows(min_row=5, min_col=ci, max_col=ci):
            if fmt:
                row[0].number_format = fmt
            if row[0].value is not None:
                longest = max(longest, len(str(row[0].value)))
        ws.column_dimensions[get_column_letter(ci)].width = min(max(longest + 2, 10), 50)
    if rep["totals"] and rep["rows"]:
        last = ws.max_row + 1
        for ci, c in enumerate(rep["columns"], 1):
            cell = ws.cell(row=last, column=ci)
            if ci == 1:
                cell.value = "TOTAL"
            elif c["key"] in rep["totals"]:
                letter = get_column_letter(ci)
                cell.value = f"=SUM({letter}5:{letter}{last - 1})"
                cell.number_format = '"R$" #,##0.00' if c["type"] == "money" else "0"
            cell.font = Font(bold=True)
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
